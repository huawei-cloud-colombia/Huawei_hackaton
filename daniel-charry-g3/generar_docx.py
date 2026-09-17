"""
generar_docx.py — Genera el documento Word del Reto 3 NEXUS LIVE.
"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

# ─── Estilos ──────────────────────────────────────────────────
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

# ─── Portada ──────────────────────────────────────────────────
title = doc.add_heading('NEXUS LIVE // T-80', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_heading('Motor de Reservas de Alta Concurrencia', level=2)
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('')
p = doc.add_paragraph('Reto Hackathon — Huawei Cloud Colombia')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

p = doc.add_paragraph('Equipo: Daniel Charry (daniel-charry-g3)')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

p = doc.add_paragraph('Fecha: 2026-09-17')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_page_break()

# ─── 1. Descripción del Reto ──────────────────────────────────
doc.add_heading('1. Descripción del Reto', level=1)
doc.add_paragraph(
    'NEXUS Live abrirá la venta general de entradas para AURORA WORLD TOUR, '
    'el evento musical más esperado del año. Más de 180.000 personas están en '
    'la sala de espera y solo hay 42.000 entradas disponibles.'
)
doc.add_paragraph(
    'El componente central de reservas (SeatLock) quedó en un estado no confiable. '
    'La misión es reconstruir un motor de reservas mínimo, correcto y verificable '
    'que garantice que un asiento nunca sea vendido a dos compradores diferentes.'
)

# ─── 2. Cómo funciona el sistema ──────────────────────────────
doc.add_heading('2. Cómo funciona el sistema', level=1)

doc.add_heading('Flujo de una reserva', level=2)
doc.add_paragraph('1. El usuario selecciona uno o varios asientos disponibles.')
doc.add_paragraph('2. Envía una solicitud POST /api/reserve con user_id, event_id y seat_ids.')
doc.add_paragraph('3. El motor valida disponibilidad (todo-o-nada), límite por usuario y crea un HOLD.')
doc.add_paragraph('4. Los asientos pasan de AVAILABLE → HELD por 120 segundos (TTL configurable).')
doc.add_paragraph('5. El usuario confirma la compra con POST /api/confirm y un payment_token.')
doc.add_paragraph('6. Si el pago es APPROVED → HELD → SOLD. La compra queda confirmada.')
doc.add_paragraph('7. Si el pago es DECLINED → los asientos se liberan (HELD → AVAILABLE).')
doc.add_paragraph('8. Si el pago es TIMEOUT/ERROR → el HOLD se mantiene para reintentar.')
doc.add_paragraph('9. Si pasan 120 segundos sin pagar → el HOLD expira automáticamente.')

doc.add_heading('Estados de un asiento', level=2)
doc.add_paragraph('🟢 AVAILABLE — Disponible para reservar', style='List Bullet')
doc.add_paragraph('🟡 HELD — Reservado temporalmente (120s)', style='List Bullet')
doc.add_paragraph('🔴 SOLD — Vendido y confirmado', style='List Bullet')

# ─── 3. Stack tecnológico ─────────────────────────────────────
doc.add_heading('3. Stack Tecnológico', level=1)

table = doc.add_table(rows=8, cols=2)
table.style = 'Light Grid Accent 1'
table.alignment = WD_TABLE_ALIGNMENT.CENTER
cells = table.rows[0].cells
cells[0].text = 'Componente'
cells[1].text = 'Tecnología'
data = [
    ('Lenguaje', 'Python 3.11+'),
    ('Framework', 'FastAPI + Uvicorn (ASGI)'),
    ('Validación', 'Pydantic v2'),
    ('Concurrencia', 'asyncio.Lock (mutex por operación atómica)'),
    ('Base de datos', 'PostgreSQL (asyncpg) + memoria'),
    ('Pagos', 'Mock con probabilidades configurables'),
    ('Resiliencia', 'Circuit Breaker (CLOSED / OPEN / HALF_OPEN)'),
]
for i, (comp, tech) in enumerate(data, 1):
    table.rows[i].cells[0].text = comp
    table.rows[i].cells[1].text = tech

# ─── 4. Arquitectura ──────────────────────────────────────────
doc.add_heading('4. Arquitectura del Proyecto', level=1)
doc.add_paragraph('Estructura de archivos:')
doc.add_paragraph(
    'daniel-charry-g3/\n'
    '├── README.md\n'
    '├── requerimientos.txt\n'
    '├── prompt_usado.txt\n'
    '├── codigo/\n'
    '│   ├── main.py          ← API FastAPI (endpoints + interfaz)\n'
    '│   ├── config.py        ← Configuración (TTL, límites, DB)\n'
    '│   ├── models.py        ← Modelos Pydantic\n'
    '│   ├── seat_lock.py     ← Motor de reservas (concurrencia + idempotencia)\n'
    '│   ├── payment.py       ← Mock de pagos + Circuit Breaker\n'
    '│   ├── audit.py         ← Trazabilidad\n'
    '│   ├── database.py      ← Esquema PostgreSQL\n'
    '│   ├── tests.py         ← Pruebas (escenarios A-F)\n'
    '│   └── static/index.html ← Interfaz web\n'
    '└── frontend/            ← Frontend Angular'
)

# ─── 5. Fases del reto ────────────────────────────────────────
doc.add_heading('5. Fases del Reto', level=1)

doc.add_heading('Fase 1 — SeatLock: Reservas Temporales', level=2)
doc.add_paragraph('• Estados AVAILABLE / HELD / SOLD', style='List Bullet')
doc.add_paragraph('• Reserva todo-o-nada (si un asiento no está disponible, falla todo)', style='List Bullet')
doc.add_paragraph('• Expiración automática (TTL configurable, default 120s)', style='List Bullet')
doc.add_paragraph('• Límite de 6 asientos por usuario', style='List Bullet')
doc.add_paragraph('• Precio controlado por el servidor', style='List Bullet')

doc.add_heading('Fase 2 — Concurrencia e Idempotencia', level=2)
doc.add_paragraph('• asyncio.Lock global para operación atómica', style='List Bullet')
doc.add_paragraph('• 100 solicitudes concurrentes sobre 1 asiento → exactamente 1 ganador', style='List Bullet')
doc.add_paragraph('• Idempotency-Key: misma clave + mismo payload = mismo resultado', style='List Bullet')
doc.add_paragraph('• Conflicto: misma clave + payload diferente = IDEMPOTENCY_CONFLICT', style='List Bullet')

doc.add_heading('Fase 3 — Confirmación y Pagos', level=2)
doc.add_paragraph('• Mock de pagos: APPROVED (70%), DECLINED (10%), ERROR (10%), TIMEOUT (10%)', style='List Bullet')
doc.add_paragraph('• APPROVED → HELD → SOLD', style='List Bullet')
doc.add_paragraph('• DECLINED → libera asientos', style='List Bullet')
doc.add_paragraph('• TIMEOUT/ERROR → mantiene HOLD para reintentar', style='List Bullet')
doc.add_paragraph('• Circuit Breaker: 3 fallos → OPEN, 15s → HALF_OPEN, 1 prueba → CLOSED/OPEN', style='List Bullet')

doc.add_heading('Fase 4 — Interfaz Visual (NEXUS Control Room)', level=2)
doc.add_paragraph('• Visualizar asientos con colores (🟢🟡🔴)', style='List Bullet')
doc.add_paragraph('• Seleccionar asientos y crear HOLD', style='List Bullet')
doc.add_paragraph('• Confirmar compra con payment_token', style='List Bullet')
doc.add_paragraph('• Simular carrera de N usuarios por 1 asiento', style='List Bullet')
doc.add_paragraph('• Ver trazabilidad y estado del Circuit Breaker', style='List Bullet')

# ─── 6. Endpoints API ─────────────────────────────────────────
doc.add_heading('6. Endpoints de la API', level=1)

table2 = doc.add_table(rows=11, cols=3)
table2.style = 'Light Grid Accent 1'
table2.rows[0].cells[0].text = 'Método'
table2.rows[0].cells[1].text = 'Endpoint'
table2.rows[0].cells[2].text = 'Función'
endpoints = [
    ('GET', '/api/seats', 'Listar asientos'),
    ('GET', '/api/seats/available', 'Asientos disponibles'),
    ('POST', '/api/reserve', 'Crear HOLD (Idempotency-Key)'),
    ('GET', '/api/holds', 'Listar reservas'),
    ('POST', '/api/holds/{id}/release', 'Liberar HOLD'),
    ('POST', '/api/confirm', 'Confirmar compra'),
    ('POST', '/api/simulate/race', 'Simular carrera'),
    ('GET', '/api/audit', 'Ver trazabilidad'),
    ('GET', '/api/payment/status', 'Estado Circuit Breaker'),
    ('POST', '/api/reset', 'Reiniciar sistema'),
]
for i, (method, endpoint, func) in enumerate(endpoints, 1):
    table2.rows[i].cells[0].text = method
    table2.rows[i].cells[1].text = endpoint
    table2.rows[i].cells[2].text = func

# ─── 7. Cómo ejecutar ─────────────────────────────────────────
doc.add_heading('7. Cómo ejecutar el sistema', level=1)
doc.add_heading('Instalación', level=2)
doc.add_paragraph('cd daniel-charry-g3/codigo')
doc.add_paragraph('python -m venv venv')
doc.add_paragraph('venv\\Scripts\\activate  (Windows)')
doc.add_paragraph('pip install -r ../requerimientos.txt')

doc.add_heading('Ejecución', level=2)
doc.add_paragraph('uvicorn main:app --host 0.0.0.0 --port 8000 --reload')
doc.add_paragraph('Abrir http://localhost:8000 en el navegador')

doc.add_heading('Pruebas', level=2)
doc.add_paragraph('pytest tests.py -v')

# ─── 8. Escenarios de prueba ──────────────────────────────────
doc.add_heading('8. Escenarios de Prueba (A-F)', level=1)
doc.add_paragraph('A — Reserva normal: AVAILABLE → HELD → SOLD', style='List Bullet')
doc.add_paragraph('B — Reserva expirada: HOLD expira → AVAILABLE', style='List Bullet')
doc.add_paragraph('C — Concurrencia: 100 usuarios → 1 ganador', style='List Bullet')
doc.add_paragraph('D — Reintento: misma Idempotency-Key = mismo resultado', style='List Bullet')
doc.add_paragraph('E — Conflicto: misma clave + payload diferente = IDEMPOTENCY_CONFLICT', style='List Bullet')
doc.add_paragraph('F — Pago degradado: Circuit Breaker OPEN, asiento no se vende', style='List Bullet')

# ─── 9. Estrategias clave ─────────────────────────────────────
doc.add_heading('9. Estrategias Clave', level=1)

doc.add_heading('Idempotencia', level=2)
doc.add_paragraph(
    'Se usa el header Idempotency-Key. Se calcula un hash SHA-256 del payload. '
    'Misma clave + mismo payload retorna el mismo resultado cacheado. '
    'Misma clave + payload diferente retorna IDEMPOTENCY_CONFLICT.'
)

doc.add_heading('Evitar Overselling', level=2)
doc.add_paragraph(
    'asyncio.Lock global durante la operación atómica de reserva. '
    'El lock protege la validación de disponibilidad y el marcado de asientos '
    'como HELD en una sola sección crítica. Garantiza 1 ganador en 100 solicitudes.'
)

doc.add_heading('Expiración de HOLDs', level=2)
doc.add_paragraph(
    'Cada HOLD tiene expires_at = created_at + TTL. '
    'Un background task ejecuta expire_holds() cada 1 segundo. '
    'Al expirar: HELD → AVAILABLE si no fue vendido.'
)

doc.add_heading('Fallos del proveedor de pagos', level=2)
doc.add_paragraph('APPROVED → HELD → SOLD (compra confirmada)')
doc.add_paragraph('DECLINED → HELD → AVAILABLE (asientos liberados)')
doc.add_paragraph('ERROR/TIMEOUT → HOLD se mantiene (no se libera, reintentar)')
doc.add_paragraph('Circuit Breaker OPEN → PAYMENT_SERVICE_UNAVAILABLE, asiento NO se marca SOLD')

# ─── 10. Entregables ──────────────────────────────────────────
doc.add_heading('10. Entregables', level=1)
doc.add_paragraph('1. Código fuente funcional', style='List Number')
doc.add_paragraph('2. README.md', style='List Number')
doc.add_paragraph('3. Archivo de dependencias (requerimientos.txt)', style='List Number')
doc.add_paragraph('4. Bitácora de uso de GLM 5.2 (prompt_usado.txt)', style='List Number')
doc.add_paragraph('5. Instrucciones de ejecución', style='List Number')
doc.add_paragraph('6. Tests (tests.py con escenarios A-F)', style='List Number')
doc.add_paragraph('7. Interfaz gráfica (HTML + Angular)', style='List Number')

# ─── Guardar ──────────────────────────────────────────────────
output_path = r'c:\Huawei_hackaton\daniel-charry-g3\NEXUS_LIVE_Reto3_Documentacion.docx'
doc.save(output_path)
print(f'Documento guardado en: {output_path}')
